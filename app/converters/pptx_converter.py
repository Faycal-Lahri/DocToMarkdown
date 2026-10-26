"""
==============================================================================
DocToMarkdown - Convertisseur PowerPoint vers Markdown (app/converters/pptx_converter.py)
==============================================================================

Ce module convertit les présentations Microsoft PowerPoint (.pptx) en Markdown :
1. Découpage diapositive par diapositive : Chaque diapositive produit une section
   claire '## Diapositive N : [Titre]' suivie d'un séparateur horizontal.
2. Extraction sémantique complète :
   - Titre de la diapositive
   - Formes de texte et sous-titres
   - Listes à puces avec gestion de l'indentation hiérarchique (- niveau 1,   - niveau 2...)
   - Tableaux PowerPoint convertis en tableaux Markdown GFM valides
   - Notes du présentateur (Speaker Notes) formatées en blocs de citation
3. Extraction d'images : Si keep_images=True, chaque image présente sur la diapositive
   est sauvegardée dans images/ et insérée avec un lien Markdown propre.
"""

from __future__ import annotations

import logging
from pathlib import Path

from app.converters.base import BaseConverter, CancellationCheck, ConversionResult, JobCancelledError, ProgressCallback

logger = logging.getLogger(__name__)


class PptxConverter(BaseConverter):
    """
    Convertisseur pour les présentations PowerPoint (.pptx) basé sur python-pptx.
    """

    def convert(
        self,
        file_path: Path,
        output_dir: Path,
        keep_images: bool = False,
        ocr: bool = False,
        progress_callback: ProgressCallback | None = None,
        cancel_check: CancellationCheck | None = None,
    ) -> ConversionResult:
        """
        Convertit une présentation .pptx en Markdown.
        """
        if cancel_check and cancel_check():
            raise JobCancelledError("Conversion PowerPoint annulée par l'utilisateur.")
        from pptx import Presentation

        output_dir.mkdir(parents=True, exist_ok=True)
        image_dir = output_dir / "images"
        if keep_images:
            image_dir.mkdir(parents=True, exist_ok=True)

        if progress_callback:
            progress_callback(5.0, 0, 0, "Chargement de la présentation PowerPoint...")

        prs = Presentation(str(file_path))
        total_slides = len(prs.slides)

        if total_slides == 0:
            return ConversionResult(
                markdown_content=f"# {file_path.stem}\n\n*Présentation PowerPoint vide (0 diapositive).*\n",
                pages_count=0,
            )

        md_output: list[str] = [f"# {file_path.stem}\n\n*Présentation PowerPoint ({total_slides} diapositives)*\n"]
        extracted_images: list[Path] = []
        warnings: list[str] = []
        global_image_counter = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_lines: list[str] = []

            # 1. Récupération du titre de la diapositive
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip().replace("\n", " ")

            if slide_title:
                slide_lines.append(f"## Diapositive {slide_idx} : {slide_title}\n")
            else:
                slide_lines.append(f"## Diapositive {slide_idx}\n")

            # 2. Parcours des formes sur la diapositive
            for shape in slide.shapes:
                # Éviter de répéter la forme du titre déjà affichée ci-dessus
                if shape == slide.shapes.title:
                    continue

                # A. Formes contenant du texte (cadres, zones, sous-titres)
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if not text:
                            continue

                        # Indentation selon le niveau de puce (level 0 = puce, level 1 = sous-puce)
                        level = para.level or 0
                        indent = "  " * level
                        slide_lines.append(f"{indent}- {text}")

                    slide_lines.append("")  # Saut de ligne après le cadre

                # B. Tableaux PowerPoint
                elif shape.has_table:
                    table_md = self._convert_table(shape.table)
                    if table_md:
                        slide_lines.append(f"\n{table_md}\n")

                # C. Images intégrées sur la diapositive
                elif keep_images and shape.shape_type == 13:  # 13 correspond à MSO_SHAPE_TYPE.PICTURE
                    try:
                        image = shape.image
                        global_image_counter += 1
                        ext = image.ext or "png"
                        img_filename = f"slide_{slide_idx}_img_{global_image_counter}.{ext}"
                        img_path = image_dir / img_filename

                        with open(img_path, "wb") as f_out:
                            f_out.write(image.blob)

                        extracted_images.append(img_path)
                        rel_link = f"images/{img_filename}"
                        slide_lines.append(f"\n![Illustration diapositive {slide_idx}]({rel_link})\n")
                    except Exception as img_err:
                        logger.warning("Impossible d'extraire une image de la slide %d: %s", slide_idx, img_err)

            # 3. Notes du présentateur (Speaker Notes)
            try:
                if slide.has_notes_slide:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        # Formatage sous forme de citation Markdown
                        notes_quote = "\n> ".join(notes_text.splitlines())
                        slide_lines.append(f"\n> **Notes du présentateur :**\n> {notes_quote}\n")
            except Exception as notes_err:
                logger.debug("Erreur lors de la lecture des notes de la slide %d: %s", slide_idx, notes_err)

            # Séparateur de diapositive Markdown
            slide_lines.append("\n---\n")
            md_output.append("\n".join(slide_lines))

            # Mise à jour du callback de progression
            if progress_callback:
                pct = 10.0 + (slide_idx / total_slides) * 85.0
                progress_callback(
                    pct, slide_idx, total_slides, f"Conversion de la diapositive {slide_idx}/{total_slides}..."
                )

        if progress_callback:
            progress_callback(100.0, total_slides, total_slides, "Présentation convertie avec succès !")

        full_markdown = "\n".join(md_output)

        return ConversionResult(
            markdown_content=full_markdown,
            image_files=extracted_images,
            warnings=warnings,
            pages_count=total_slides,
        )

    def _convert_table(self, table) -> str:
        """
        Convertit un objet tableau PowerPoint en tableau Markdown format GFM.
        """
        rows = table.rows
        if not rows:
            return ""

        table_lines = []

        # En-têtes (première ligne)
        header_cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in rows[0].cells]
        table_lines.append("| " + " | ".join(header_cells) + " |")

        # Séparateurs
        separators = ["---" for _ in header_cells]
        table_lines.append("| " + " | ".join(separators) + " |")

        # Lignes suivantes
        for row in rows[1:]:
            row_cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
            table_lines.append("| " + " | ".join(row_cells) + " |")

        return "\n".join(table_lines)
